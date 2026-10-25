"""The outline as a Word document, and the same document read back with writing in it.

A grant professional who tested GrantBridge asked to draft in Word and upload the result. The round trip has to
tell the funder's requirements apart from the writer's prose, and it must never quietly lose a draft.
"""

import io

import pytest

from grantbridge import auth, outline, outline_docx as od, repo
from tests.test_pages import PASSWORD

SECTIONS = [
    {"heading": "Need", "purpose": "Why the acquisition is needed.", "limit_words": 600, "limit_pages": None,
     "elements": [{"text": "Describe the community you serve."},
                  {"text": "Quantify the need with data."}],
     "evidence": [{"text": "An aerial map of the property."}], "scored_by": ["Criterion 3"], "response": ""},
    {"heading": "Approach", "purpose": "", "limit_words": None, "limit_pages": 2,
     "elements": [{"text": "Explain how the objectives will be attained."}],
     "evidence": [], "scored_by": [], "response": "Our existing approach text."},
]
PAYLOAD = {"components": [{"name": "SF-424", "condition": "Required from all applicants."}],
           "registrations": [{"name": "SAM.gov", "lead_time": "can take several months"}],
           "format_rules": [{"text": "Closing Date for Applications: 10/31/2026"}]}


def built(sections=None, payload=None) -> bytes:
    return od.build("Highlands 2027", "U.S. Fish and Wildlife Service",
                    SECTIONS if sections is None else sections,
                    PAYLOAD if payload is None else payload)


def write_under(data: bytes, heading: str, text: str) -> bytes:
    """Edit the file the way a person would: replace everything written under one heading.

    The body of a section is every paragraph between its heading and the next one, so this clears them all and
    writes into the first. Setting only the paragraph after the heading would hit the blank spacer that
    follows the boxed note and leave the real draft untouched.
    """
    import docx

    document = docx.Document(io.BytesIO(data))
    paragraphs = document.paragraphs
    start = None
    for i, paragraph in enumerate(paragraphs):
        if paragraph.style.name.startswith("Heading 1") and paragraph.text.strip() == heading:
            start = i
            break
    if start is None:
        raise AssertionError(f"no heading {heading!r} in the document")
    end = len(paragraphs)
    for j in range(start + 1, len(paragraphs)):
        if paragraphs[j].style.name.startswith("Heading"):
            end = j
            break
    for j in range(start + 1, end):
        paragraphs[j].text = ""
    paragraphs[start + 1].text = text
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- what goes out

def test_the_document_carries_every_section_in_the_funders_order():
    import docx

    document = docx.Document(io.BytesIO(built()))
    headings = [p.text.strip() for p in document.paragraphs if p.style.name.startswith("Heading 1")]
    assert headings[:2] == ["Need", "Approach"]


def test_the_requirements_travel_with_each_section():
    import docx

    document = docx.Document(io.BytesIO(built()))
    boxed = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "Describe the community you serve." in boxed
    assert "Quantify the need with data." in boxed
    assert "Attach or show: An aerial map of the property." in boxed
    assert "Limit: 600 words." in boxed and "Limit: 2 pages." in boxed
    assert "Judged against: Criterion 3" in boxed


def test_writing_already_in_grantbridge_is_in_the_file():
    import docx

    document = docx.Document(io.BytesIO(built()))
    body = "\n".join(p.text for p in document.paragraphs)
    assert "Our existing approach text." in body


def test_the_forms_and_registrations_are_listed_but_are_not_sections_to_write():
    import docx

    document = docx.Document(io.BytesIO(built()))
    body = "\n".join(p.text for p in document.paragraphs)
    assert "SF-424 — Required from all applicants." in body
    assert "SAM.gov — can take several months" in body
    drafts, problems = od.read_back(built(), SECTIONS)
    assert [d.heading for d in drafts] == ["Need", "Approach"]   # the appendix is not one of them
    assert problems == []


# ---------------------------------------------------------------- what comes back

def test_an_untouched_file_changes_nothing():
    drafts, problems = od.read_back(built(), SECTIONS)
    assert problems == []
    assert all(not d.changed for d in drafts)
    assert od.summarise(drafts, SECTIONS)["changed"] == 0


def test_writing_added_in_word_comes_back():
    data = write_under(built(), "Need", "We serve 1,200 households across two counties.")
    drafts, _ = od.read_back(data, SECTIONS)
    need = next(d for d in drafts if d.heading == "Need")
    assert need.changed and need.text == "We serve 1,200 households across two counties."
    assert od.summarise(drafts, SECTIONS)["new_text"] == 1


def test_the_boxed_requirements_never_come_back_as_the_writers_prose():
    """The whole reason the guidance sits in a table."""
    drafts, _ = od.read_back(built(), SECTIONS)
    need = next(d for d in drafts if d.heading == "Need")
    assert "Describe the community you serve." not in need.text
    assert "Why the acquisition is needed." not in need.text      # the purpose line, too
    assert need.text == ""


def test_a_renamed_heading_is_reported_and_nothing_is_lost():
    import docx

    document = docx.Document(io.BytesIO(built()))
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading 1") and paragraph.text.strip() == "Need":
            paragraph.text = "Statement of Need"
    buffer = io.BytesIO()
    document.save(buffer)
    drafts, problems = od.read_back(buffer.getvalue(), SECTIONS)
    assert [d.heading for d in drafts] == ["Approach"]
    assert any("was not found in the document" in p and "Need" in p for p in problems)


def test_text_under_an_unknown_heading_is_reported_rather_than_dropped_silently():
    import docx

    document = docx.Document(io.BytesIO(built()))
    document.add_heading("Letters of Support", level=1)
    document.add_paragraph("Three letters are attached to this application for review.")
    buffer = io.BytesIO()
    document.save(buffer)
    _, problems = od.read_back(buffer.getvalue(), SECTIONS)
    assert any("Letters of Support" in p and "not a section of this outline" in p for p in problems)


def test_emptying_a_section_is_counted_so_it_can_be_warned_about():
    sections = [dict(SECTIONS[1])]                       # the one that already has writing
    data = write_under(od.build("T", "F", sections, {}), "Approach", "")
    drafts, _ = od.read_back(data, sections)
    counts = od.summarise(drafts, sections)
    assert counts["emptied"] == 1 and counts["changed"] == 1


def test_a_file_that_is_not_a_word_document_is_refused_by_its_bytes():
    with pytest.raises(ValueError):
        od.read_back(b"%PDF-1.7 pretending", SECTIONS)


# ---------------------------------------------------------------- writing it back into the application

@pytest.fixture
def application():
    user = auth.sign_up("alice@example.org", PASSWORD, "Alice Org")
    app_id = repo.create_application(user.id, "Highlands 2027")
    outline.apply_import(user.id, app_id, {
        "sections": [{**s, "elements": [{**e, "excerpt": "x", "source": "nofo.txt"} for e in s["elements"]],
                      "evidence": [], "excerpt": "x"} for s in SECTIONS],
        "criteria": [], "components": [], "registrations": [], "format_rules": []}, {})
    return user, app_id


def test_confirmed_drafts_are_saved_and_the_status_moves_on(application):
    user, app_id = application
    stored = repo.list_outline_sections(user.id, app_id)
    data = write_under(od.build("Highlands 2027", "FWS", stored, {}), "Need",
                       "We serve 1,200 households across two counties.")
    drafts, _ = od.read_back(data, stored)
    assert od.apply(user.id, app_id, drafts, repo) == 1
    after = {s["heading"]: s for s in repo.list_outline_sections(user.id, app_id)}
    assert after["Need"]["response"] == "We serve 1,200 households across two counties."
    assert after["Need"]["status"] == "drafting"          # it had not been started before


def test_sections_the_file_did_not_change_are_left_alone(application):
    user, app_id = application
    stored = repo.list_outline_sections(user.id, app_id)
    repo.save_outline_response(user.id, stored[1]["id"], "Approach written in the browser.", status="ready")
    stored = repo.list_outline_sections(user.id, app_id)
    drafts, _ = od.read_back(od.build("T", "F", stored, {}), stored)
    assert od.apply(user.id, app_id, drafts, repo) == 0
    after = repo.list_outline_sections(user.id, app_id)[1]
    assert after["response"] == "Approach written in the browser." and after["status"] == "ready"
