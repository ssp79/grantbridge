"""The proposal outline as a Word document, and the same document read back with the writing in it.

Plenty of grant writers would rather draft in Word than in a browser, and plenty of organizations circulate a
draft for comment before anyone types it anywhere. So the outline goes out as a .docx laid out in the funder's
own order, and comes back with whatever was written under each heading.

How the round trip stays honest:

  - Each section is a Word **Heading 1** carrying the funder's own wording. Matching on the way back is by
    heading text, normalized for case and spacing.
  - The funder's requirements sit in a single-cell **table** under each heading, not in the body text. Word
    shows them as a boxed note; python-docx keeps table text out of `document.paragraphs`. So when the file
    comes back, every body paragraph under a heading is the writer's own prose and nothing else — the
    instructions cannot be mistaken for the draft.
  - A heading that was renamed, deleted or added is **reported, not guessed at**. GrantBridge says which
    sections it could not place and leaves those drafts alone.
  - Nothing is saved until the user confirms the summary of what would change.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

MAX_DOCX_BYTES = 8 * 1024 * 1024
GUIDANCE_HEADING = "What this section must contain"
APPENDIX = "Also required by this funder"   # GrantBridge's own page, never read back
INTRO = (
    "This outline was read out of the funder's own solicitation. Write your draft as ordinary paragraphs under "
    "each heading, below the boxed note.\n\n"
    "Keep the headings exactly as they are and leave them styled as Heading 1. GrantBridge matches your work "
    "to its sections by heading when you upload this file back, and it will tell you about any heading it "
    "cannot place rather than guessing. The boxed notes are the funder's requirements: leave them alone, or "
    "delete them if you prefer — they are ignored when the file is read back."
)


@dataclass(frozen=True)
class SectionDraft:
    heading: str
    text: str
    matched: bool          # a section GrantBridge holds
    changed: bool          # different from what is stored
    words: int


def _norm(text: str) -> str:
    return " ".join((text or "").split()).strip().lower().rstrip(":")


# ---------------------------------------------------------------- out to Word

def build(title: str, funder: str, sections: list[dict], payload: dict | None = None) -> bytes:
    """The outline as a .docx, in the funder's order, with any draft already written."""
    import docx
    from docx.shared import Pt

    document = docx.Document()
    document.add_heading(title or "Grant application", level=0)
    meta = [bit for bit in (funder, _deadline_of(payload)) if bit]
    if meta:
        document.add_paragraph(" · ".join(meta))
    note = document.add_paragraph(INTRO)
    note.runs[0].italic = True

    for section in sections:
        document.add_page_break()
        document.add_heading(section["heading"], level=1)
        _guidance_box(document, section, Pt)
        body = (section.get("response") or "").strip()
        for paragraph in (body.split("\n") if body else [""]):
            document.add_paragraph(paragraph)

    extras = _extras(payload)
    if extras:
        document.add_page_break()
        document.add_heading(APPENDIX, level=1)
        document.add_paragraph("These are not sections you write here. They are listed so nothing is "
                               "forgotten; GrantBridge ignores this page when the file is read back.")
        for heading, lines in extras:
            document.add_heading(heading, level=2)
            for line in lines:
                document.add_paragraph(line, style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _deadline_of(payload: dict | None) -> str:
    for rule in (payload or {}).get("format_rules") or []:
        if "due" in rule.get("text", "").lower() or "closing date" in rule.get("text", "").lower():
            return rule["text"]
    return ""


def _guidance_box(document, section: dict, Pt) -> None:
    """The funder's requirements, in a one-cell table so they stay out of the body text."""
    lines = []
    if section.get("purpose"):
        lines.append(section["purpose"])
    limit = section.get("limit_pages") or section.get("limit_words")
    if section.get("limit_pages"):
        lines.append(f"Limit: {section['limit_pages']} pages.")
    elif section.get("limit_words"):
        lines.append(f"Limit: {section['limit_words']} words.")
    for element in section.get("elements") or []:
        lines.append(f"• {element['text']}")
    for element in section.get("evidence") or []:
        lines.append(f"• Attach or show: {element['text']}")
    if section.get("scored_by"):
        lines.append("Judged against: " + ", ".join(section["scored_by"]))
    if not lines:
        return
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    cell.text = GUIDANCE_HEADING
    cell.paragraphs[0].runs[0].bold = True
    for line in lines:
        paragraph = cell.add_paragraph(line)
        for run in paragraph.runs:
            run.font.size = Pt(9)
    document.add_paragraph("")
    _ = limit


def _extras(payload: dict | None) -> list[tuple[str, list[str]]]:
    payload = payload or {}
    out = []
    components = payload.get("components") or []
    if components:
        out.append(("Forms and attachments",
                    [f"{c['name']} — {c['condition']}" for c in components]))
    registrations = payload.get("registrations") or []
    if registrations:
        out.append(("Registrations",
                    [f"{r['name']}" + (f" — {r['lead_time']}" if r.get("lead_time") else "")
                     for r in registrations]))
    rules = payload.get("format_rules") or []
    if rules:
        out.append(("Format and submission rules", [r["text"] for r in rules]))
    return out


# ---------------------------------------------------------------- back from Word

def read_back(data: bytes, sections: list[dict]) -> tuple[list[SectionDraft], list[str]]:
    """Pull each heading's prose out of an edited copy. Returns (drafts, problems).

    Only body paragraphs count, so the boxed requirements never come back as if the writer had typed them.
    """
    if len(data) > MAX_DOCX_BYTES:
        raise ValueError(f"That file is larger than {MAX_DOCX_BYTES // 1_048_576} MB.")
    # The same rule as every other upload: the file's own bytes decide what it is, not its name.
    from grantbridge import documents as docs

    try:
        kind = docs.sniff(data, "outline.docx")
    except docs.UploadError as err:
        raise ValueError(str(err)) from None
    if kind != "docx":
        raise ValueError("That is not a Word .docx file. Download the outline again, edit that copy, and "
                         "upload it back.")
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as err:
        raise ValueError("That Word document could not be read. Save it again as .docx and try once more.") \
            from err

    found: dict[str, list[str]] = {}
    order: list[str] = []
    current: str | None = None
    for paragraph in document.paragraphs:
        style = (paragraph.style.name or "") if paragraph.style is not None else ""
        text = paragraph.text.strip()
        if style.startswith("Heading 1"):
            if _norm(text) == _norm(APPENDIX):
                current = None       # GrantBridge wrote this page; it is not anyone's draft
                continue
            current = text
            if current and current not in found:
                found[current] = []
                order.append(current)
            continue
        if style.startswith("Heading") or style == "Title":
            continue                      # the document title, or the "Also required" page's subheadings
        if current and text:
            found[current].append(text)

    known = {_norm(s["heading"]): s for s in sections}
    drafts, problems = [], []
    for heading in order:
        section = known.get(_norm(heading))
        body = "\n\n".join(found[heading]).strip()
        if section is None:
            if body:
                problems.append(f"“{heading}” is not a section of this outline, so its {len(body.split())} "
                                "word(s) were left out. Rename the heading to match, or add the section first.")
            continue
        stored = (section.get("response") or "").strip()
        drafts.append(SectionDraft(heading=section["heading"], text=body, matched=True,
                                   changed=body != stored, words=len(body.split())))
    seen = {_norm(h) for h in order}
    for section in sections:
        if _norm(section["heading"]) not in seen:
            problems.append(f"“{section['heading']}” was not found in the document, so what you have here was "
                            "left as it is.")
    return drafts, problems


def summarise(drafts: list[SectionDraft], sections: list[dict]) -> dict:
    """What confirming would do, counted so it can be shown before anything is written."""
    stored = {_norm(s["heading"]): (s.get("response") or "").strip() for s in sections}
    changed = [d for d in drafts if d.changed]
    return {
        "sections_in_file": len(drafts),
        "changed": len(changed),
        "unchanged": len(drafts) - len(changed),
        "new_text": sum(1 for d in changed if not stored.get(_norm(d.heading))),
        "replaced": sum(1 for d in changed if stored.get(_norm(d.heading))),
        "emptied": sum(1 for d in changed if not d.text and stored.get(_norm(d.heading))),
        "words": sum(d.words for d in drafts),
    }


def apply(user_id: int, application_id: int, drafts: list[SectionDraft], repo) -> int:
    """Write the confirmed drafts back. Returns how many sections changed."""
    by_heading = {_norm(s["heading"]): s for s in repo.list_outline_sections(user_id, application_id)}
    written = 0
    for draft in drafts:
        if not draft.changed:
            continue
        section = by_heading.get(_norm(draft.heading))
        if section is None:
            continue
        status = section["status"] if not draft.text else (
            "drafting" if section["status"] == "not started" else section["status"])
        repo.save_outline_response(user_id, section["id"], draft.text, status=status)
        written += 1
    return written
